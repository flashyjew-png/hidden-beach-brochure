#!/usr/bin/env python3
"""Build sheet-template/hidden-beach-brochure.xlsx: the starter Google Sheet for the brochure.

Jake uploads the .xlsx to Google Drive and saves it as a Google Sheet (see README.md).
Tabs and header rows are the renderer's content contract; test/sheet-template.test.js checks
them against fixtures/*.csv and the Settings keys the code reads.

Every cell is written as text (format '@') so Google Sheets keeps prices, phone numbers and
hours exactly as typed instead of turning them into numbers, dates or times.

Run from anywhere:  python3 sheet-template/build.py   (needs openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import TYPE_STRING
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parent / "hidden-beach-brochure.xlsx"

# Extra empty rows pre-formatted as text, so new rows Jake types stay text too.
TEXT_ROWS = 200

SETTINGS_HEADER = ["key", "en", "th", "de"]
SETTINGS = [
    ("business_name", "Hidden Beach"),
    ("logo_url", ""),
    ("hero_photo", ""),
    ("status_banner", "Beach bar & restaurant at Hidden Beach Resort, Koh Mak"),
    ("about", "Beachfront bar at Hidden Beach Resort on Koh Mak, right on our own private beach. "
              "Come by for a cold drink with your toes in the sand."),
    ("heading_about", "About us"),
    ("heading_cocktails", "Cocktails"),
    ("heading_food", "Food"),
    ("food_coming_soon", "New menu coming soon"),
    ("activities", "Make a day of it at Hidden Beach Resort: our private beach, kayaking and scuba diving. "
                   "Ask at the bar."),
    ("maps_url", "https://www.google.com/maps/search/?api=1&query=Hidden+Beach+Resort+Koh+Mak"),
    ("rating", "4.8★ from 192 Google reviews"),
    ("whatsapp_number", "+1 907 215 8419"),
    ("whatsapp_greeting", "Hi! I found you via your brochure"),
    ("heading_specials", "Specials & events"),
    ("heading_activities", "Make a day of it"),
    ("heading_find_us", "Find us"),
    ("heading_hours", "Opening hours"),
    ("label_open_today", "Open today"),
    ("label_closed_today", "Closed today"),
    ("label_closed", "Closed"),
    ("label_day", "Day"),
    ("label_bar", "Bar"),
    ("label_kitchen", "Kitchen"),
    ("label_maps_button", "Open in Google Maps"),
    ("label_whatsapp_button", "Chat with us on WhatsApp"),
    ("label_language", "Language"),
]

HOURS_HEADER = ["day", "bar", "kitchen"]
HOURS = [[day, "10:00–23:00", "11:00–17:00"] for day in ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]]

MENU_HEADER = ["menu", "category", "name_en", "name_th", "name_de", "desc_en", "desc_th", "desc_de",
               "price", "photo", "tags", "show"]


def menu_row(menu, category, name, desc, tags, show):
    # Prices left blank on purpose: Jake fills in the real ones.
    return [menu, category, name, "", "", desc, "", "", "", "", tags, show]


MENU = [
    menu_row("cocktails", "Classics", "Mojito", "White rum, lime, fresh mint, soda", "", "yes"),
    menu_row("cocktails", "Classics", "Piña Colada", "Rum, pineapple, coconut cream", "", "yes"),
    menu_row("cocktails", "Classics", "Mai Tai", "Rum, lime, orange liqueur, orgeat", "", "yes"),
    menu_row("food", "Mains", "Green curry", "Example dish – replace with your own", "spicy", "no"),
    menu_row("food", "Mains", "Pad Thai", "Example dish – replace with your own", "", "no"),
    menu_row("food", "Snacks", "Spring rolls", "Example dish – replace with your own", "vegetarian", "no"),
]

SPECIALS_HEADER = ["title_en", "title_th", "title_de", "detail_en", "detail_th", "detail_de",
                   "when", "photo", "show"]
SPECIALS = [
    ["Sunset happy hour", "", "", "Example – change the details, then set show to yes", "", "", "Daily 5–7pm", "", "no"],
    ["Live music", "", "", "Example – change the details, then set show to yes", "", "", "Saturday from 8pm", "", "no"],
]

# (tab name, header, rows, column widths)
TABS = [
    ("Settings", SETTINGS_HEADER, [[k, en, "", ""] for k, en in SETTINGS], [22, 60, 40, 40]),
    ("Hours", HOURS_HEADER, HOURS, [10, 16, 16]),
    ("Menu", MENU_HEADER, MENU, [11, 14, 22, 18, 18, 36, 24, 24, 9, 30, 14, 7]),
    ("Specials", SPECIALS_HEADER, SPECIALS, [22, 18, 18, 40, 24, 24, 18, 30, 7]),
]


def text_cell(ws, row, col, value):
    cell = ws.cell(row=row, column=col)
    cell.number_format = "@"
    if value != "":
        cell.value = value
        cell.data_type = TYPE_STRING  # explicit: never a number/date
    return cell


def build(path=OUT):
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows, widths in TABS:
        ws = wb.create_sheet(name)
        for c, h in enumerate(header, 1):
            text_cell(ws, 1, c, h).font = Font(bold=True)
        for r, values in enumerate(rows, 2):
            for c, v in enumerate(values, 1):
                text_cell(ws, r, c, v).alignment = Alignment(wrap_text=True, vertical="top")
        for r in range(len(rows) + 2, TEXT_ROWS + 2):
            for c in range(1, len(header) + 1):
                text_cell(ws, r, c, "")
        ws.freeze_panes = "A2"
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"Wrote {build()}")
